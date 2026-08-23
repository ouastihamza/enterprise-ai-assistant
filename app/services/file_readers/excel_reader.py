from pathlib import Path

from openpyxl import load_workbook

from app.services.file_readers.base_reader import BaseFileReader


class ExcelReader(BaseFileReader):
    """
    Extracts readable text from Excel .xlsx files.
    """

    def extract_text(self, file_path: str) -> str:
        excel_path = Path(file_path)

        if not excel_path.exists():
            raise FileNotFoundError(f"Excel file not found: {file_path}")

        workbook = load_workbook(
            filename=excel_path,
            data_only=True,
        )

        extracted_rows = []

        for sheet in workbook.worksheets:
            extracted_rows.append(f"Sheet: {sheet.title}")

            rows = list(sheet.iter_rows(values_only=True))

            if not rows:
                continue

            headers = rows[0]

            for row_number, row in enumerate(rows[1:], start=2):
                values = []

                for header, value in zip(headers, row):
                    if header is None and value is None:
                        continue

                    column_name = str(header) if header is not None else "Unknown Column"
                    cell_value = str(value) if value is not None else ""

                    values.append(f"{column_name}: {cell_value}")

                if values:
                    extracted_rows.append(
                        f"Row {row_number}: " + " | ".join(values)
                    )

        return "\n".join(extracted_rows)