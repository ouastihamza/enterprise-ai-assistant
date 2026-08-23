from pathlib import Path

from pptx import Presentation

from app.services.file_readers.base_reader import BaseFileReader


class PPTXReader(BaseFileReader):
    """
    Extracts readable text from PowerPoint presentations.
    """

    def extract_text(self, file_path: str) -> str:

        ppt_path = Path(file_path)

        if not ppt_path.exists():
            raise FileNotFoundError(
                f"PPTX file not found: {file_path}"
            )

        presentation = Presentation(ppt_path)

        extracted_text = []

        for slide_number, slide in enumerate(
            presentation.slides,
            start=1,
        ):
            extracted_text.append(
                f"Slide {slide_number}"
            )

            for shape in slide.shapes:

                if hasattr(shape, "text"):

                    text = shape.text.strip()

                    if text:
                        extracted_text.append(text)

        return "\n".join(extracted_text)